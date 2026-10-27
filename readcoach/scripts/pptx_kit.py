"""把 答辩幻灯片.html 重建成真正的 .pptx（文字可编辑，不是截图）。

坐标不另起一套：HTML 的画布是 1280×720px，PPT 用 13.333×7.5in，
两者恰好 1px = 9525 EMU，所以下面所有数字都能和 HTML 里的数值一一对上，
要改某一块的位置，直接照 HTML 的 px 抄。

两张矢量图（封面插画、架构图）走图片：它们是示意图形，用形状硬搬
既费事又容易在 PowerPoint 里变形。其余全是文本框、圆角矩形和真表格，
字号用 px×0.75 换算（96dpi 下 1pt = 1.333px）。
"""

from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

# ---------- 基本换算 ----------
W, H = 1280, 720
EMU_PER_PX = 9525


def px(v: float) -> Emu:
    return Emu(int(round(v * EMU_PER_PX)))


def fs(v: float) -> Pt:
    """CSS px → pt"""
    return Pt(round(v * 0.75, 1))


# ---------- 调色板（与 HTML 的 CSS 变量一致） ----------
BG = RGBColor(0xF6, 0xF8, 0xFC)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x0F, 0x17, 0x2A)
BODY = RGBColor(0x25, 0x32, 0x4A)
MUTED = RGBColor(0x5B, 0x6B, 0x84)
LINE = RGBColor(0xDF, 0xE6, 0xF2)
ACCENT = RGBColor(0x25, 0x63, 0xEB)
ACCENT_SOFT = RGBColor(0xE8, 0xEF, 0xFD)
ACCENT_DEEP = RGBColor(0x1D, 0x4E, 0xD8)
TEAL = RGBColor(0x0F, 0x9D, 0x8F)
TEAL_SOFT = RGBColor(0xE3, 0xF6, 0xF3)
TEAL_DEEP = RGBColor(0x0B, 0x7A, 0x70)
WARN = RGBColor(0xD9, 0x77, 0x06)
WARN_SOFT = RGBColor(0xFE, 0xF3, 0xE2)
WARN_DEEP = RGBColor(0xB4, 0x53, 0x09)
DANGER = RGBColor(0xDC, 0x26, 0x26)
DANGER_SOFT = RGBColor(0xFD, 0xEA, 0xEA)
DANGER_DEEP = RGBColor(0xB9, 0x1C, 0x1C)
OK = RGBColor(0x16, 0xA3, 0x4A)
OK_SOFT = RGBColor(0xE8, 0xF7, 0xEE)
OK_DEEP = RGBColor(0x15, 0x80, 0x3D)
PURPLE = RGBColor(0x7C, 0x3A, 0xED)

CJK = "微软雅黑"
MONO = "Consolas"


# ---------- 文字 ----------
TOKEN = re.compile(r"(\*\*[^*]+\*\*|__[^_]+__|`[^`]+`)")


def _font(run, size, bold=False, color=INK, mono=False):
    f = run.font
    f.name = MONO if mono else CJK
    f.size = fs(size)
    f.bold = bold
    f.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = rpr.makeelement(qn("a:ea"), {})
        rpr.append(ea)
    ea.set("typeface", CJK)


def _fill_runs(p, text, size, color, bold, mono, highlight=None):
    """支持 **粗体**、__强调色__、`代码` 三种行内标记。"""
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run()
            r.text = part[2:-2]
            _font(r, size, True, color, mono)
        elif part.startswith("__") and part.endswith("__"):
            r = p.add_run()
            r.text = part[2:-2]
            _font(r, size, True, ACCENT, mono)
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run()
            r.text = part[1:-1]
            _font(r, size * 0.94, bold, RGBColor(0x1E, 0x3A, 0x8A), False)
        else:
            r = p.add_run()
            r.text = part
            _font(r, size, bold, color, mono)


def text(slide, x, y, w, h, content, size=19, color=BODY, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, mono=False,
         spacing=1.35, space_after=0, margins=(0, 0, 0, 0)):
    """content 可以是字符串，也可以是段落列表（每段一个字符串）。"""
    box = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    tf.margin_left, tf.margin_right, tf.margin_top, tf.margin_bottom = [px(m) for m in margins]

    paras = content if isinstance(content, list) else [content]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        if space_after:
            p.space_after = px(space_after)
        _fill_runs(p, para, size, color, bold, mono)
    return box


def _no_shadow(shape):
    sp_pr = shape._element.spPr
    if sp_pr.find(qn("a:effectLst")) is None:
        sp_pr.append(sp_pr.makeelement(qn("a:effectLst"), {}))


def rect(slide, x, y, w, h, fill=CARD, line=LINE, radius=0, line_w=1.2, shape=None):
    kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    s = slide.shapes.add_shape(kind, px(x), px(y), px(w), px(h))
    if radius:
        # pptx 的圆角值是「半径 / 短边」
        s.adjustments[0] = min(0.5, radius / min(w, h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    _no_shadow(s)
    s.text_frame.word_wrap = True
    return s


def card(slide, x, y, w, h, fill=CARD, line=LINE, radius=12, bar=None):
    rect(slide, x, y, w, h, fill=fill, line=line, radius=radius)
    if bar:
        rect(slide, x + 1, y + 10, 4, h - 20, fill=bar, line=None, radius=2)


def set_grad(shape, stops, ang=5400000):
    """给形状换成线性渐变填充。stops = [(位置0~1, RGB), ...]，ang 单位 1/60000 度。"""
    sp_pr = shape._element.spPr
    for tag in ("a:solidFill", "a:gradFill", "a:noFill"):
        old = sp_pr.find(qn(tag))
        if old is not None:
            sp_pr.remove(old)
    grad = sp_pr.makeelement(qn("a:gradFill"), {"rotWithShape": "1"})
    lst = grad.makeelement(qn("a:gsLst"), {})
    for pos, color in stops:
        gs = lst.makeelement(qn("a:gs"), {"pos": str(int(pos * 100000))})
        clr = gs.makeelement(qn("a:srgbClr"), {"val": "%02X%02X%02X" % (color[0], color[1], color[2])})
        gs.append(clr)
        lst.append(gs)
    grad.append(lst)
    lin = grad.makeelement(qn("a:lin"), {"ang": str(ang), "scaled": "0"})
    grad.append(lin)
    # 必须按 schema 顺序插在 a:ln 之前，否则整段 fill 会被 PowerPoint 丢掉（页面变全白）
    ln = sp_pr.find(qn("a:ln"))
    if ln is not None:
        ln.addprevious(grad)
    else:
        sp_pr.append(grad)


def bg_grad(slide, stops, ang=5400000):
    s = rect(slide, 0, 0, 1280, 720, fill=BG, line=None)
    set_grad(s, stops, ang)
    slide.shapes._spTree.remove(s._element)
    slide.shapes._spTree.insert(2, s._element)
    return s


def badge(slide, x, y, n, color=ACCENT):
    s = rect(slide, x, y, 28, 28, fill=color, line=None, radius=8)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _fill_runs(p, str(n), 15, RGBColor(0xFF, 0xFF, 0xFF), True, False)


def header(slide, kicker, title, rule_color=ACCENT, y_rule=190):
    text(slide, 68, 50, 1000, 24, kicker, size=14, color=ACCENT, bold=True)
    text(slide, 68, 78, 1144, 60, title, size=38, color=INK, bold=True, spacing=1.2)
    rect(slide, 68, y_rule, 64, 4, fill=rule_color, line=None, radius=2)


def pageno(slide, n, total=12):
    text(slide, 1060, 682, 152, 20, f"{n:02d} / {total}", size=13,
         color=RGBColor(0x9A, 0xA8, 0xBF), align=PP_ALIGN.RIGHT)


# ---------- 代码块 ----------
KW = {"def", "return", "if", "not", "in", "for", "is", "else", "or", "None", "str",
      "raise", "import", "from", "class", "with", "as", "and", "True", "False"}
CODE_COLORS = {
    "p": RGBColor(0xE2, 0xE8, 0xF5),
    "c": RGBColor(0x8F, 0xA3, 0xC4),
    "k": RGBColor(0x7D, 0xD3, 0xFC),
    "s": RGBColor(0xA7, 0xF3, 0xD0),
    "e": RGBColor(0xFC, 0xA5, 0xA5),
}


def _tokens(line: str):
    """极简高亮：注释 / 字符串 / 关键字，够复现 HTML 里的那几段。"""
    if line.strip().startswith("#"):
        return [(line, "c")]
    out, i, buf = [], 0, ""
    while i < len(line):
        ch = line[i]
        if ch == '"':
            if buf:
                out.append((buf, "p"))
                buf = ""
            j = line.find('"', i + 1)
            j = len(line) - 1 if j < 0 else j
            out.append((line[i:j + 1], "s"))
            i = j + 1
        elif ch == "#" and line[max(0, i - 1)] != "\\":
            if buf:
                out.append((buf, "p"))
                buf = ""
            out.append((line[i:], "c"))
            i = len(line)
        else:
            buf += ch
            i += 1
    if buf:
        out.append((buf, "p"))
    # 关键字：在非注释片段里按词切
    final = []
    for seg, kind in out:
        if kind != "p":
            final.append((seg, kind))
            continue
        for piece in re.split(r"(\b\w+\b)", seg):
            if piece in KW:
                final.append((piece, "k"))
            elif piece:
                final.append((piece, "p"))
    return final


def code(slide, x, y, w, h, lines, size=14.5, bg=RGBColor(0x0F, 0x17, 0x2A), pad=18,
         spacing=1.55):
    rect(slide, x, y, w, h, fill=bg, line=None, radius=12)
    box = slide.shapes.add_textbox(px(x + pad + 2), px(y + pad), px(w - pad * 2 - 4), px(h - pad * 2))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = spacing
        segs = line if isinstance(line, list) else _tokens(line)
        for seg, kind in segs:
            r = p.add_run()
            r.text = seg
            _font(r, size, False, CODE_COLORS[kind], mono=True)


# ---------- 表格 ----------
def _cell_line(cell, color=LINE, width=0.75):
    tc_pr = cell._tc.get_or_add_tcPr()
    for side in ("lnB",):
        tag = tc_pr.find(qn(f"a:{side}"))
        if tag is not None:
            tc_pr.remove(tag)
        ln = tc_pr.makeelement(qn(f"a:{side}"), {"w": str(int(width * 12700)), "cap": "flat",
                                                 "cmpd": "sng", "algn": "ctr"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        clr = fill.makeelement(qn("a:srgbClr"), {"val": "%02X%02X%02X" % (color[0], color[1], color[2])})
        fill.append(clr)
        ln.append(fill)
        tc_pr.append(ln)


def _cell_fill(cell, color):
    cell.fill.solid()
    cell.fill.fore_color.rgb = color


def table(slide, x, y, w, rows, col_w, size=16, head_size=13.5, row_h=44, head_h=40,
          col_fill=None, pad_v=8):
    """rows[0] 是表头；每一格可以是字符串（支持行内标记）。
    col_fill: {列号: 颜色}，用来给「本项目」那一列加底色。"""
    n_rows, n_cols = len(rows), len(rows[0])
    shape = slide.shapes.add_table(n_rows, n_cols, px(x), px(y), px(w),
                                   px(head_h + row_h * (n_rows - 1)))
    tbl = shape.table
    tbl.firstRow = False
    tbl.horzBanding = False
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = px(cw)
    tbl.rows[0].height = px(head_h)
    for r in range(1, n_rows):
        tbl.rows[r].height = px(row_h)

    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = px(12)
            cell.margin_top = cell.margin_bottom = px(pad_v)
            cell.vertical_anchor = MSO_ANCHOR.TOP
            _cell_fill(cell, (col_fill or {}).get(c, CARD) if r else CARD)
            _cell_line(cell)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.line_spacing = 1.3
            if r == 0:
                _fill_runs(p, val, head_size, MUTED, True, False)
            else:
                _fill_runs(p, val, size, BODY, False, False)
    return shape


def text_width(s: str, size: float) -> float:
    """粗估一行字有多宽：中日韩按 1em、拉丁按 0.55em。"""
    return sum(1.0 if ord(ch) > 0x2E80 else 0.55 for ch in s) * size


def pill(slide, x, y, label, fg=ACCENT_DEEP, bg=ACCENT_SOFT, size=13.5, pad=12):
    w = max(56, text_width(label, size) + pad * 2)
    s = rect(slide, x, y, w, 26, fill=bg, line=None, radius=13)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _fill_runs(p, label, size, fg, True, False)
    return x + w + 6
