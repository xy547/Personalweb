#!/usr/bin/env python
"""在《答辩幻灯片.pptx》里插入三页「demo 亮点」，并把讲稿放到页面上。

三页分别是：
    8  关键能力 · 多语种     —— 五门语言 + 「查不到就明说」的边界
    9  关键能力 · 每日推送   —— 三级回落 + 连续打卡
    10 关键能力 · 学习闭环   —— 词表 → Word 笔记 → 练习页 → 错词本 → 间隔复习

插入位置：原第 7 页「权限隔离」之后（能力展示连成一段），随后再回到「与课堂实验的关系」。

讲稿的放法：每页底部一条浅色横条，左边一个「讲稿」标签，右边是照念的正文；
同一段话也写进 PowerPoint 备注栏（演讲者视图能直接看）。
字号比正文小、底色比正文淡，所以不抢视觉重心。

用法：
    python scripts/add_highlight_pages.py 源文件.pptx [输出文件.pptx]
    python scripts/add_highlight_pages.py 源文件.pptx --inplace
"""
from __future__ import annotations

import re
import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from pptx_kit import (ACCENT, ACCENT_DEEP, ACCENT_SOFT, BG, BODY, CARD, INK, LINE,
                      MUTED, OK, OK_DEEP, OK_SOFT, PURPLE, TEAL, TEAL_DEEP, TEAL_SOFT,
                      WARN, WARN_DEEP, WARN_SOFT, _fill_runs, card, header, pageno,
                      pill, px, rect, text)

# 语种卡片配色
LANGS = [
    ("EN", "英语", "English", "CET-4/6 · 考研", "IPA 音标", "6161 条可核验", ACCENT, ACCENT_SOFT),
    ("JA", "日语", "日本語", "JLPT N5–N1", "假名 + 罗马音", "无本地词库", TEAL, TEAL_SOFT),
    ("FR", "法语", "Français", "DELF · DALF", "IPA 音标", "无本地词库", PURPLE, RGBColor(0xF1, 0xEB, 0xFE)),
    ("DE", "德语", "Deutsch", "Goethe A1–C2", "IPA 音标", "无本地词库", WARN, WARN_SOFT),
    ("ES", "西语", "Español", "DELE A1–C2", "IPA 音标", "无本地词库", OK, OK_SOFT),
]


def bg(slide, color=BG):
    s = rect(slide, 0, 0, 1280, 720, fill=color, line=None, radius=0)
    slide.shapes._spTree.remove(s._element)
    slide.shapes._spTree.insert(2, s._element)


# ---------------- 讲稿条 ----------------
def script_bar(slide, y, lines, color=ACCENT, soft=ACCENT_SOFT, label="讲稿"):
    """页脚上方的讲稿条：浅底 + 竖线 + 标签 + 照念正文（不抢正文重心）。"""
    h = 104
    card(slide, 68, y, 1144, h, fill=soft, line=None, radius=10)
    rect(slide, 70, y + 10, 3, h - 20, fill=color, line=None, radius=2)
    pill(slide, 88, y + 12, label, fg=color, bg=RGBColor(0xFF, 0xFF, 0xFF), size=12.5, pad=11)
    text(slide, 88, y + 46, 1104, h - 52, lines, size=13, color=RGBColor(0x33, 0x3D, 0x4B),
         spacing=1.45)


# ---------------- 新增第 8 页：多语种 ----------------
def s_lang(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    header(slide, "关键能力 · 多语种", "五门语言，一套流程，__不假装能查__")

    x0, w, gap, y0, h = 68, 214, 16, 212, 150
    for i, (abbr, cn, en, exam, phon, lib, color, soft) in enumerate(LANGS):
        x = x0 + i * (w + gap)
        card(slide, x, y0, w, h, bar=color)
        badge_rect = rect(slide, x + 16, y0 + 16, 34, 34, fill=color, line=None, radius=9)
        tf = badge_rect.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = 3  # MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        _fill_runs(p, abbr, 13, RGBColor(0xFF, 0xFF, 0xFF), True, False)
        text(slide, x + 60, y0 + 18, w - 74, 22, cn, size=17, color=INK, bold=True)
        text(slide, x + 60, y0 + 40, w - 74, 18, en, size=12, color=MUTED)
        rows = [("考纲", exam), ("注音", phon), ("词库", lib)]
        for j, (k, v) in enumerate(rows):
            yy = y0 + 66 + j * 26
            text(slide, x + 16, yy, 34, 20, k, size=12, color=MUTED)
            text(slide, x + 50, yy, w - 66, 22, v, size=12.5,
                 color=INK if j < 2 else (OK_DEEP if lib.startswith("6161") else WARN_DEEP),
                 bold=(j == 2))

    card(slide, 68, 378, 556, 148, bar=ACCENT)
    text(slide, 92, 394, 508, 26, "考纲跟着语种走，不是一套模板套五种语言",
         size=17, color=INK, bold=True)
    text(slide, 92, 428, 508, 86,
         "左侧切到日语，「目标等级」立刻从 `CET-4/6/考研` 换成 `JLPT N5–N1`；"
         "法语是 `DELF/DALF`，德语 `Goethe`，西语 `DELE`。\n"
         "级别体系、注音写法、词形规矩都按各语种的考试规矩来——德语名词要带 der/die/das。",
         size=13.5, color=BODY, spacing=1.5)

    card(slide, 644, 378, 568, 148, bar=WARN)
    text(slide, 668, 394, 520, 26, "查不到就明说查不到，不给假考纲等级",
         size=17, color=INK, bold=True)
    text(slide, 668, 428, 520, 86,
         "非英语没有本地词库，查词工具直接回 `NOT_SUPPORTED`（**不是** `NOT_FOUND`），"
         "模型就不会以为「查过但没有」而反复重试。\n"
         "这类词表里每条都标「模型推断」，Word 顶部写明「本语种暂无本地词库」。",
         size=14.5, color=BODY, spacing=1.5)

    script_bar(slide, 540, [
        "这一页是我做完主路径之后最想说的一个升级：原来只能学英语，现在左侧选日语、法语、德语、西班牙语，"
        "等级列表会跟着换——切到日语就是 JLPT N5 到 N1，不是继续写六级。",
        "但边界我要讲清楚：非英语我没有做本地词库，所以查词工具会直接说「这个语种不支持」，"
        "词表里每条都标「模型推断」。我不给它编一个假的考纲等级，宁可让它说不行。",
    ])
    pageno(slide, 8, 15)
    return slide


# ---------------- 新增第 9 页：每日推送 ----------------
def s_daily(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    header(slide, "关键能力 · 每日推送", "每天一篇，__自己送到门口__")

    card(slide, 68, 212, 700, 316, bar=ACCENT)
    text(slide, 92, 228, 652, 26, "一天一篇，三级回落", size=18, color=INK, bold=True)
    text(slide, 92, 258, 652, 22, "不是定时任务在后台跑，是用户打开就有内容",
         size=13, color=MUTED)

    steps = [
        ("1", "当天缓存", "同一语种当天只抓一次，重复打开不重复抓网。", ACCENT, ACCENT_SOFT),
        ("2", "联网抓取", "去新华网各语种版取当天正文（英/日/法/德/西 各有站点）。", TEAL, TEAL_SOFT),
        ("3", "离线快照", "抓不到就用 `data/articles/` 里的存档，断网也能用。", WARN, WARN_SOFT),
    ]
    for i, (no, name, desc, color, soft) in enumerate(steps):
        y = 292 + i * 78
        card(slide, 92, y, 652, 66, fill=soft, line=None, radius=10)
        b = rect(slide, 108, y + 19, 28, 28, fill=color, line=None, radius=8)
        tf = b.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = 3
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        _fill_runs(p, no, 13, RGBColor(0xFF, 0xFF, 0xFF), True, False)
        text(slide, 150, y + 12, 200, 24, name, size=16, color=INK, bold=True)
        text(slide, 150, y + 38, 570, 24, desc, size=13.5, color=BODY)

    card(slide, 788, 212, 424, 208, bar=OK)
    text(slide, 812, 228, 376, 26, "今天（10-27）真的推了五篇", size=17, color=INK, bold=True)
    rows = [("🇬🇧 英语", "Economic Watch · China's…", "3626 字"),
            ("🇯🇵 日语", "中国スマホ・アフリカの…", "1606 字"),
            ("🇫🇷 法语", "Chine : paysages d'automne", "1299 字"),
            ("🇩🇪 德语", "Vorbereitungen für 9. CIIE…", "2078 字"),
            ("🇪🇸 西语", "Principal provincia…", "2262 字")]
    for i, (lang, title, size) in enumerate(rows):
        y = 264 + i * 30
        text(slide, 812, y, 96, 22, lang, size=13, color=INK, bold=True)
        text(slide, 906, y, 200, 22, title, size=11, color=MUTED)
        text(slide, 1110, y, 78, 22, size, size=12, color=OK_DEEP, bold=True)

    card(slide, 788, 436, 424, 92, bar=TEAL)
    text(slide, 812, 452, 376, 24, "连续打卡按日期集合算", size=16, color=INK, bold=True)
    text(slide, 812, 480, 376, 40,
         "漏一天就断，不是计数器；首页直接显示连续天数。",
         size=13.5, color=BODY, spacing=1.45)

    script_bar(slide, 540, [
        "第二个亮点是每天推送。左边这三级回落是实际跑出来的：先看当天缓存，没有就联网去新华网各语种版抓，"
        "抓不到就用离线快照兜底，所以断网也有内容。",
        "右边是今天的真实结果，五个语种各抓到一篇，右下角是连续打卡天数——按日期集合算的，漏一天就断。",
    ])
    pageno(slide, 9, 15)
    return slide


# ---------------- 新增第 10 页：学习闭环 ----------------
def s_loop(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    header(slide, "关键能力 · 学习闭环", "词表只是起点，__闭环才留得住__")

    steps = [("1", "挑词", "从原文挑生词"), ("2", "CSV 词表", "导进背单词软件"),
             ("3", "Word 笔记", "排好版直接打印"), ("4", "练习页", "原文自动挖空"),
             ("5", "错词本", "做错的自动记下"), ("6", "间隔复习", "到期才出现")]
    w, gap, x0, y0, h = 178, 15, 68, 212, 92
    for i, (no, name, desc) in enumerate(steps):
        x = x0 + i * (w + gap)
        card(slide, x, y0, w, h, bar=ACCENT if i < 4 else TEAL)
        text(slide, x + 16, y0 + 14, 40, 22, no, size=15, color=MUTED, bold=True)
        text(slide, x + 16, y0 + 36, w - 32, 24, name, size=15, color=INK, bold=True)
        text(slide, x + 16, y0 + 62, w - 32, 22, desc, size=12, color=MUTED)
        if i < len(steps) - 1:
            text(slide, x + w - 2, y0 + 34, 19, 24, "→", size=15, color=RGBColor(0xB4, 0xC0, 0xD2),
                 align=PP_ALIGN.CENTER)

    card(slide, 68, 326, 556, 200, bar=OK)
    text(slide, 92, 342, 508, 26, "今天真的产出了这些文件", size=17, color=INK, bold=True)
    files = [("n3_daily_vocab_20261027.docx", "日语 N3 词表笔记（Word）"),
             ("wordlist.csv", "五列词表，逗号分隔，可导 Excel"),
             ("campus_carbon_worksheet.html", "原文挖空的练习页"),
             ("learner_profile.json", "错词本与复习日期，落成文件")]
    for i, (f, d) in enumerate(files):
        y = 380 + i * 34
        text(slide, 92, y, 300, 22, f, size=13, color=ACCENT_DEEP, mono=True)
        text(slide, 396, y, 204, 22, d, size=12.5, color=MUTED)

    card(slide, 644, 326, 568, 200, bar=TEAL)
    text(slide, 668, 342, 520, 26, "复习日期由代码算，不由模型算", size=17, color=INK, bold=True)
    text(slide, 668, 378, 520, 134,
         "错词写进 `data/learner_profile.json`，每次记录都带 `first_seen / last_seen / ease / interval`；"
         "`due_words` 只取**今天到期**的词，明天到期的今天不出现。\n"
         "模型只负责「哪个词错了」，**什么时候复习**是代码说了算。",
         size=14.5, color=BODY, spacing=1.55)

    script_bar(slide, 540, [
        "最后是闭环。词表不是终点：挑完词，CSV 给你导进背单词软件，Word 笔记是排好版能直接打印的，"
        "练习页把原文挖空，做错的词自动进错词本，复习日期由代码按间隔算。",
        "左边这些是我今天真的生成的文件。所以第 2 页说的第四步——挖空练习——现在是自动的，而且做完还能回头复习。",
    ], color=TEAL, soft=TEAL_SOFT)
    pageno(slide, 10, 15)
    return slide


# ---------------- 现场演示页：补第 6 条话术 + 修学号占位 ----------------
def add_demo_extra(prs):
    """演示页原本五条话术排到 572px，没有空间塞第六张大卡。
    改成在页脚左侧放一行低调的彩蛋提示，不动原有五条版式。"""
    target = None
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.has_text_frame and "五条话术" in sh.text_frame.text:
                target = slide
                break
        if target is not None:
            break
    if target is None:
        return

    text(target, 68, 680, 720, 20,
         "彩蛋（时间富余时）：切到日语看首页今日推送 —— 对应第 8 页",
         size=12, color=RGBColor(0x9A, 0xA8, 0xBF))

    notes = target.notes_slide.notes_text_frame
    notes.text = notes.text.rstrip() + (
        "\n6）彩蛋：切到日语，首页今日推送那篇，然后说「开始今天的精读」。"
        "等级自动变 JLPT N3，词表带假名注音，每条标「模型推断」。")


def fix_cover_note(prs):
    """封面备注里还写着「学号 ×××」，换成真实学号。"""
    for slide in prs.slides:
        if not slide.has_notes_slide:
            continue
        tf = slide.notes_slide.notes_text_frame
        if "学号" not in tf.text:
            continue
        for para in tf.paragraphs:
            for run in para.runs:
                if "×××" in run.text:
                    run.text = run.text.replace("×××", "f25016515")
        return


# ---------------- 页码重排 + 插入 ----------------
def renumber(prs, total):
    pat = re.compile(r"^\d{2} / \d{2}$")
    for i, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            t = sh.text_frame.text.strip()
            if not pat.match(t):
                continue
            runs = [r for p in sh.text_frame.paragraphs for r in p.runs]
            if not runs:
                continue
            runs[0].text = f"{i:02d} / {total}"
            for r in runs[1:]:
                r.text = ""


def move_slides(prs, new_ids, target_index):
    lst = prs.slides._sldIdLst
    for sid in new_ids:
        lst.remove(sid)
    for i, sid in enumerate(new_ids):
        lst.insert(target_index + i, sid)


def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    src = Path(args[0]) if args else Path(r"C:\Users\xuyue\Desktop\答辩幻灯片.pptx")
    inplace = "--inplace" in sys.argv
    dst = Path(args[1]) if len(args) > 1 else (src if inplace else src.with_name(src.stem + "-含亮点.pptx"))

    prs = Presentation(str(src))
    before = len(prs.slides)

    s_lang(prs)
    s_daily(prs)
    s_loop(prs)

    set_notes(prs.slides[-3], SCRIPT_LANG)
    set_notes(prs.slides[-2], SCRIPT_DAILY)
    set_notes(prs.slides[-1], SCRIPT_LOOP)

    total = before + 3

    add_demo_extra(prs)      # 演示页补第 6 条话术
    fix_cover_note(prs)      # 封面备注里的学号占位

    ids = list(prs.slides._sldIdLst)
    move_slides(prs, ids[-3:], 7)   # 插到第 8 页的位置（原第 8 页之前）
    renumber(prs, total)            # 必须在移动之后，否则页码按旧顺序算

    prs.save(str(dst))
    print(f"原 {before} 页 -> 现 {total} 页，已插入到第 8/9/10 页")
    print(f"输出：{dst}")


SCRIPT_LANG = """这一页是我做完主路径之后最想说的一个升级：原来只能学英语，现在左侧选日语、法语、德语、西班牙语，等级列表会跟着换——切到日语就是 JLPT N5 到 N1，不是继续写六级。

但边界我要讲清楚：非英语我没有做本地词库，所以查词工具会直接说「这个语种不支持」，词表里每条都标「模型推断」。我不给它编一个假的考纲等级，宁可让它说不行。"""

SCRIPT_DAILY = """第二个亮点是每天推送。左边这三级回落是实际跑出来的：先看当天缓存，没有就联网去新华网各语种版抓，抓不到就用离线快照兜底，所以断网也有内容。

右边是今天的真实结果，五个语种各抓到一篇，右下角是连续打卡天数——按日期集合算的，漏一天就断。"""

SCRIPT_LOOP = """最后是闭环。词表不是终点：挑完词，CSV 给你导进背单词软件，Word 笔记是排好版能直接打印的，练习页把原文挖空，做错的词自动进错词本，复习日期由代码按间隔算。

左边这些是我今天真的生成的文件。所以第 2 页说的第四步——挖空练习——现在是自动的，而且做完还能回头复习。"""

if __name__ == "__main__":
    main()
