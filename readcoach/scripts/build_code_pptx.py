"""把 代码解读.html 重建成可编辑的 .pptx（13 页，16:9）。

坐标沿用 HTML 的 1280×720 画布（1px = 9525 EMU），所以下面每个数字
都能和 HTML 里的 padding / font-size 对上，改版式不用另算一套。

内容不手抄：走 deck_parse 解析原 HTML，代码块连高亮色一起搬过来，
所以 PPT 与网页版逐字一致。首页、地图页、收束页结构特殊，单独铺。
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from deck_parse import parse  # noqa: E402
import pptx_kit as kit  # noqa: E402
from pptx_kit import (  # noqa: E402
    ACCENT, ACCENT_DEEP, ACCENT_SOFT, BG, BODY, CARD, CJK, INK, LINE, MONO, MUTED,
    OK_DEEP, OK_SOFT, TEAL_DEEP, TEAL_SOFT, WARN, WARN_DEEP, WARN_SOFT,
    _fill_runs, _font, bg_grad, card, code, fs, pill, px, rect, text, text_width,
)

ROOT = HERE.parent
OUT = ROOT / "代码解读.pptx"

ROLE = RGBColor(0x1E, 0x3A, 0x8A)   # 行内 code 的深蓝
PAGE_BG = RGBColor(0xF6, 0xF8, 0xFC)

kit.CODE_COLORS["n"] = RGBColor(0xFC, 0xD3, 0x4D)   # 名字（金色）

TONE = {
    "plain": ACCENT,
    "teal": RGBColor(0x0B, 0x7A, 0x70),
    "warn": WARN,
    "danger": RGBColor(0xB9, 0x1C, 0x1C),
}

M = 62                    # 左右留白，与 .slide{padding:40px 62px 58px} 一致
CW = 1280 - M * 2         # 内容宽 1156
BOTTOM = 720 - 58         # 内容底边 662
GAP = 16

# ---------- 量字：估行数，决定字号与卡片高度 ----------
def _em(ch: str, size: float) -> float:
    return size * (1.0 if ord(ch) > 0x2E80 else 0.55)


def count_lines(runs, width: float, size: float) -> int:
    """runs 里 code 段按 0.92 倍字号算；返回折行后的行数。"""
    lines, cur = 1, 0.0
    for text_, kind in runs:
        s = size * (0.92 if kind == "code" else 1.0)
        for ch in text_:
            if ch == "\n":
                lines, cur = lines + 1, 0.0
                continue
            w = _em(ch, s)
            if cur + w > width:
                lines, cur = lines + 1, w
            else:
                cur += w
    return lines


def count_code_lines(line, width: float, size: float) -> int:
    """等宽：拉丁 0.55em，中日韩 1em。"""
    n, cur = 1, 0.0
    for seg, _cls in line:
        for ch in seg:
            w = size * (1.0 if ord(ch) > 0x2E80 else 0.55)
            if cur + w > width:
                n, cur = n + 1, w
            else:
                cur += w
    return n


def fit_code(lines, w: float, h: float, pad: float = 14) -> float:
    inner_w, inner_h = w - pad * 2 - 4, h - pad * 2
    for size in (12.6, 12.2, 11.8, 11.4, 11.0, 10.6, 10.2, 9.8, 9.4, 9.0):
        rows = sum(count_code_lines(ln, inner_w, size) for ln in lines)
        # 行距设成固定值（见 code_block），所以这里的乘法就是真实高度
        if rows * size * 1.66 * 1.03 <= inner_h:
            return size
    return 9.0


def code_block(slide, x, y, w, h, lines, size, pad=14, line=1.66):
    """比 pptx_kit.code() 多一件事：行距用固定 Pt。

    用倍数行距时 PowerPoint 会按字体自带的 1.17em 再乘一次，
    23 行的块实际比 CSS 高出近 20%，最后一屏就直接掉出深色底。
    固定 Pt 之后，rows × size × 1.66 就是真实高度，排版可预算。
    """
    rect(slide, x, y, w, h, fill=RGBColor(0x0F, 0x17, 0x2A), line=None, radius=11)
    box = slide.shapes.add_textbox(px(x + pad), px(y + pad), px(w - pad * 2), px(h - pad * 2))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    step = Pt(size * line * 0.75)
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = step
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        for seg, kind in ln:
            r = p.add_run()
            r.text = seg
            _font(r, size, False, kit.CODE_COLORS[kind], mono=True)
    return box


# ---------- 行内富文本（批注正文 / 收束语） ----------
def rich(slide, x, y, w, h, runs, size, color, spacing=1.5, bold=False,
         b_color=None, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.line_spacing = Pt(size * spacing * 0.75)   # 固定行高，估算和实际才对得上
    for t, kind in runs:
        if kind == "nl":
            p = tf.add_paragraph()
            p.line_spacing = Pt(size * spacing * 0.75)
            continue
        r = p.add_run()
        r.text = t
        if kind == "code":
            _font(r, size * 0.92, bold, ROLE, mono=True)
        elif kind in ("b", "b2"):
            _font(r, size, True, INK if kind == "b2" else (b_color or color), False)
        else:
            _font(r, size, bold, color, False)
    return box


def kicker_title(slide, pg):
    text(slide, M, 36, CW, 20, pg["kicker"], size=13, color=ACCENT, bold=True, spacing=1.3)
    text(slide, M, 58, CW, 40, pg["h2"], size=27, color=INK, bold=True, spacing=1.2)


# ---------- 3~12 页：左代码、右批注、底部一句话 ----------
def slide_code(prs, pg):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, 1280, 720, fill=PAGE_BG, line=None)
    kicker_title(s, pg)

    tk = pg.get("takeaway", [])
    tk_len = len("".join(t for t, _ in tk))
    tk_rows = 1 if tk_len <= 74 else 2
    tk_h = 18 + 21 * tk_rows + 8
    tk_y = BOTTOM - tk_h

    gy = 112
    gh = tk_y - 14 - gy
    lw = round((CW - GAP) * 0.59)
    rw = CW - GAP - lw

    need = []
    for n in pg.get("notes", []):
        body_lines = count_lines(n["body"], rw - 24, 13.6)
        title_lines = count_lines(n["title"], rw - 24, 11.5)
        need.append(18 + 17 * title_lines + 3 + body_lines * 20.4)
    need_h = sum(need) + 9 * max(0, len(need) - 1)

    # 两列共用一个高度：谁的内容长就按谁算，底边对齐。
    # 深色块刚好包住代码，不会像固定满高那样在下面吊一截空底色。
    if pg.get("code"):
        size = fit_code(pg["code"], lw, gh)
        rows = sum(count_code_lines(ln, lw - 32, size) for ln in pg["code"])
        code_h = rows * size * 1.66 * 1.03 + 28
    else:
        size, code_h = 12.6, 0
    col_h = min(gh, max(code_h, need_h))

    if pg.get("code"):
        code_block(s, M, gy, lw, col_h, pg["code"], size=size)

    nx = M + lw + GAP
    gap_n = 9
    if need and col_h > need_h:
        spare = col_h - need_h
        need = [h + spare / len(need) for h in need]

    y = gy
    for i, (n, h) in enumerate(zip(pg.get("notes", []), need)):
        card(s, nx, y, rw, h, radius=9, line=LINE)
        t_lines = count_lines(n["title"], rw - 24, 11.5)
        rich(s, nx + 12, y + 9, rw - 24, 17 * t_lines, n["title"], 11.5,
             TONE[n["tone"]], spacing=1.3, bold=True, b_color=TONE[n["tone"]])
        rich(s, nx + 12, y + 9 + 17 * t_lines + 3, rw - 24, h - 27,
             n["body"], 13.6, RGBColor(0x2B, 0x3A, 0x52), spacing=1.5)
        y += h + (gap_n if i < len(need) - 1 else 0)
    if y > gy + gh + 2:
        print(f"  ! 第 {pg['no']} 页批注超出可用区 {round(y - gy - gh)}px")

    if tk:
        rect(s, M, tk_y, CW, tk_h, fill=ACCENT_SOFT, line=RGBColor(0xC9, 0xDC, 0xFF), radius=9)
        rich(s, M + 14, tk_y + 9, CW - 28, tk_h - 18, tk, 14.2, BODY,
             spacing=1.5, b_color=ACCENT_DEEP)

    pageno(s, pg["no"])
    return s


def pageno(s, n):
    text(s, 1000, 688, 218, 18, f"{n:02d} / 13", size=12.5,
         color=RGBColor(0x9A, 0xA8, 0xBF), align=PP_ALIGN.RIGHT)


# ---------- 01 封面 ----------
def slide_cover(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg_grad(s, [(0.0, (0xEE, 0xF3, 0xFD)), (0.42, (0xF7, 0xF9, 0xFD)),
                (0.80, (0xF4, 0xFA, 0xF9)), (1.0, (0xE7, 0xF6, 0xF3))], 2700000)
    text(s, M, 182, CW, 20, "期末项目 · 代码拆解", size=13, color=ACCENT, bold=True, spacing=1.3)
    text(s, M, 206, CW, 130, ["十段代码，", "十次「不把正确性交给模型」"],
         size=52, color=INK, bold=True, spacing=1.16)
    rect(s, M, 366, 58, 4, fill=ACCENT, line=None, radius=2)
    rich(s, M, 392, 900, 72, [
        ("外刊精读教练（ReadCoach）· 每段都标了真实文件与行号，都回答三个问题：", "p"),
        ("\n", "nl"),
        ("为什么这么写", "b2"), ("、", "p"), ("逻辑怎么走", "b2"), ("、", "p"),
        ("用到哪些知识", "b2"), ("。", "p"),
    ], 20, RGBColor(0x33, 0x41, 0x5C), spacing=1.6, b_color=INK)

    x = M
    for label, fg, bg in (
        ("配置层 ×1", ACCENT_DEEP, ACCENT_SOFT),
        ("编排层 ×2", TEAL_DEEP, TEAL_SOFT),
        ("工具层 ×6", WARN_DEEP, WARN_SOFT),
        ("界面层 ×1", OK_DEEP, OK_SOFT),
    ):
        x = pill(s, x, 492, label, fg=fg, bg=bg, size=13, pad=11) + 2
    text(s, M, 542, CW, 20, "下方箭头或滚轮翻页 · 按 F 全屏 · 共 13 页", size=13.5, color=MUTED)
    pageno(s, 1)
    return s


# ---------- 02 地图 ----------
def flow_row(s, x, y, items, size=14):
    for item in items:
        if item == "→":
            text(s, x, y + 5, 18, 24, "→", size=16, color=RGBColor(0x9A, 0xA8, 0xBF),
                 bold=True, align=PP_ALIGN.CENTER)
            x += 18 + 8
            continue
        w = text_width(item, size) + 24
        box = rect(s, x, y, w, 34, fill=CARD, line=LINE, radius=9)
        tf = box.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        _fill_runs(p, item, size, INK, True, False)
        x += w + 8


def layer_tag(s, x, y, label, color):
    w = text_width(label, 12) + 18
    box = rect(s, x, y, w, 22, fill=color, line=None, radius=5)
    tf = box.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _fill_runs(p, label, 12, RGBColor(0xFF, 0xFF, 0xFF), True, False)
    return w


def slide_map(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, 1280, 720, fill=PAGE_BG, line=None)
    text(s, M, 36, CW, 20, "地图", size=13, color=ACCENT, bold=True, spacing=1.3)
    text(s, M, 58, CW, 40, "十段代码分布在四层，一条调用链串起来", size=27, color=INK,
         bold=True, spacing=1.2)
    rect(s, M, 108, 58, 4, fill=ACCENT, line=None, radius=2)

    flow_row(s, M, 136, [
        "用户一句话", "→", "app.py 界面", "→", "build_system_prompt", "→",
        "StateGraph 循环", "→", "七个工具", "→", "outputs/ 文件",
    ])

    col_w = (CW - GAP * 2) / 3
    cards = [
        (RGBColor(0x64, 0x74, 0x8B), "配置层", "写盘权限边界",
         [("① ", "safe_output_path()", "config.py:40")],
         "Agent 只能写 outputs/，越界直接报错。"),
        (RGBColor(0x0F, 0x9D, 0x8F), "编排层", "循环与提示词",
         [("② ", "StateGraph", "graph.py:43"),
          ("③ ", "build_system_prompt()", "prompts.py:126")],
         "agent ⇄ tools 循环；技能按需注入。"),
        (RGBColor(0xD9, 0x77, 0x06), "工具层", "接口、数据、校验",
         [("④ ", "@tool + docstring", "dictionary.py:22"),
          ("⑤ ", "load_lexicon()", "lexicon.py:17"),
          ("⑥ ", "_verify()", "wordlist.py:68"),
          ("⑦ ", "_drop_source_field()", "wordlist.py:35"),
          ("⑧ ", "_word_pattern()", "worksheet.py:22"),
          ("⑨ ", "record_mistakes()", "memory.py:135")],
         "七个工具共用一份词库、一套校验。"),
    ]
    top, h = 190, 392
    for i, (color, tag, title, rows, note) in enumerate(cards):
        x = M + i * (col_w + GAP)
        card(s, x, top, col_w, h, radius=12)
        layer_tag(s, x + 19, top + 16, tag, color)
        text(s, x + 19, top + 47, col_w - 38, 24, title, size=16, color=INK, bold=True)
        y = top + 78
        for num, name, loc in rows:
            rich(s, x + 19, y, col_w - 38, 44, [
                (num, "p"), (name, "code"), ("\n", "nl"), (loc, "p"),
            ], 15, RGBColor(0x25, 0x32, 0x4A), spacing=1.45)
            y += 42
        text(s, x + 19, y + 6, col_w - 38, 40, note, size=13.5, color=MUTED, spacing=1.4)

    y = top + h + 22
    layer_tag(s, M, y, "界面层", RGBColor(0x16, 0xA3, 0x4A))
    rich(s, M + 78, y - 2, CW - 78, 28, [
        ("⑩ ", "p"), ("graph.stream()", "code"), (" 事件循环", "p"),
        ("　　app.py:105 —— 把每一步摊给用户看", "p"),
    ], 15, RGBColor(0x25, 0x32, 0x4A), spacing=1.4)

    pageno(s, 2)
    return s


# ---------- 13 收束 ----------
def slide_end(prs, pg):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, 1280, 720, fill=PAGE_BG, line=None)
    text(s, M, 36, CW, 20, pg["kicker"] or "收束", size=13, color=ACCENT, bold=True, spacing=1.3)
    text(s, M, 58, CW, 40, pg["h2"] or "十段代码，共用一条判断", size=27, color=INK,
         bold=True, spacing=1.2)
    rect(s, M, 108, 58, 4, fill=ACCENT, line=None, radius=2)

    col_w = (CW - GAP) / 2
    top, h = 140, 300

    card(s, M, top, col_w, h, radius=12)
    text(s, M + 19, top + 16, col_w - 38, 24, "按知识类型归类", size=16, color=INK, bold=True)
    rows = [
        ("语言基础", "装饰器、生成器、类型注解、functools、pathlib"),
        ("数据处理", "JSON 持久化、CSV 编码（utf-8-sig）、正则与词边界"),
        ("工程习惯", "不可信输入、权限最小化、fail fast、单一数据源、缓存"),
        ("Agent 层", "状态图、Function Calling、上下文分层、按需加载、记忆"),
    ]
    y = top + 52
    for i, (label, body) in enumerate(rows):
        text(s, M + 19, y + 2, 84, 24, label, size=14, color=INK, bold=True)
        text(s, M + 19 + 84, y, col_w - 38 - 84, 40, body, size=15,
             color=RGBColor(0x25, 0x32, 0x4A), spacing=1.45)
        y += 56
        if i < 3:
            rect(s, M + 12, y - 10, col_w - 24, 1, fill=LINE, line=None)

    rx = M + col_w + GAP
    card(s, rx, top, col_w, h, fill=ACCENT_SOFT, line=RGBColor(0xC9, 0xDC, 0xFF), radius=12)
    text(s, rx + 19, top + 16, col_w - 38, 24, "如果重做，我会改三处", size=16, color=INK,
         bold=True)
    fixes = [
        ("① 词库换真实词典 API", "　6161 条考试词表够撑住大部分，但派生词（retrieval / coherence）"
                              "仍查不到，只能退回「模型推断 + 标注」。"),
        ("② 工具参数用 JSON 而非竖线", "　竖线省了模型输出，却引出了「六列挤位」那个 bug。"),
        ("③ 把七个工具挂成 MCP server", "　参数表和返回格式都不用改，只是换个暴露方式。"),
    ]
    y = top + 52
    for head, body in fixes:
        rich(s, rx + 16, y, col_w - 26, 68, [(head, "b"), (body, "p")], 13.8,
             RGBColor(0x25, 0x32, 0x4A), spacing=1.55, b_color=INK)
        y += 76

    x = M
    tags = [("① 路径", ACCENT_DEEP, ACCENT_SOFT), ("② 循环", ACCENT_DEEP, ACCENT_SOFT),
            ("③ 上下文", TEAL_DEEP, TEAL_SOFT), ("④ 接口", TEAL_DEEP, TEAL_SOFT),
            ("⑤ 一致性", WARN_DEEP, WARN_SOFT), ("⑥ 校验", WARN_DEEP, WARN_SOFT),
            ("⑦ 容错", WARN_DEEP, WARN_SOFT), ("⑧ 算法", WARN_DEEP, WARN_SOFT),
            ("⑨ 记忆", WARN_DEEP, WARN_SOFT), ("⑩ 可见性", OK_DEEP, OK_SOFT)]
    for label, fg, bg in tags:
        x = pill(s, x, 466, label, fg=fg, bg=bg, size=13, pad=11)

    if pg.get("takeaway"):
        rect(s, M, 508, CW, 78, fill=ACCENT_SOFT, line=RGBColor(0xC9, 0xDC, 0xFF), radius=9)
        rich(s, M + 14, 517, CW - 28, 60, pg["takeaway"], 15, BODY, spacing=1.5,
             b_color=ACCENT_DEEP)

    pageno(s, 13)
    return s


def main() -> None:
    pages = parse()
    prs = Presentation()
    prs.slide_width = Emu(1280 * 9525)
    prs.slide_height = Emu(720 * 9525)

    for pg in pages:
        if pg["no"] == 1:
            slide_cover(prs)
        elif pg["no"] == 2:
            slide_map(prs)
        elif pg["no"] == 13:
            slide_end(prs, pg)
        else:
            slide_code(prs, pg)

    prs.save(OUT)
    print(f"OK {OUT}  {OUT.stat().st_size / 1024:.0f} KB  {len(prs.slides._sldIdLst)} 页")


if __name__ == "__main__":
    main()
