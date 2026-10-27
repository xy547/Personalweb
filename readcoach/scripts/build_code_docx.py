#!/usr/bin/env python
"""把《代码解读》(13 页 pptx) 的内容排成 Word 讲义。

数据来源 `_code_deck.json`（由 scripts/deck_parse.py 从 代码解读.html 解析而来），
所以 Word 里的代码与幻灯片逐字一致；改 HTML 后重跑 deck_parse + 本脚本即可同步。

用法：
    python scripts/build_code_docx.py            # -> 代码解读.docx
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
from md2docx import (  # noqa: E402  复用同一套字体/底纹/边框工具，保证两份 Word 版式统一
    ACCENT,
    CJK,
    INK,
    MONO,
    MUTED,
    _keep_with_next,
    _para_border,
    _set_font,
    _shade,
)

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "_code_deck.json"
OUT = ROOT / "代码解读.docx"

# 语法高亮配色（浅底上的可读版本，与网页深色底不同）
CODE_COLORS = {
    "k": (RGBColor(0x1B, 0x4F, 0xC4), True),   # 关键字
    "s": (RGBColor(0xA1, 0x3A, 0x2E), False),  # 字符串
    "n": (RGBColor(0x6D, 0x33, 0x9C), False),  # 名字 / 函数
    "c": (RGBColor(0x74, 0x7F, 0x8C), False),  # 注释
    "p": (RGBColor(0x2A, 0x34, 0x44), False),  # 普通
}
TONE_COLORS = {
    "plain": (RGBColor(0x1F, 0x29, 0x37), "D8DEE9"),
    "teal": (RGBColor(0x0B, 0x6B, 0x63), "9FD3CB"),
    "warn": (RGBColor(0x9A, 0x4A, 0x06), "F0C08A"),
}


def spans_text(spans) -> str:
    return "".join(s[0] for s in spans)


def add_spans(p, spans, size=10.5, color=INK, bold=False):
    """把 [[文本, 类型], ...] 写进段落；类型 code 用等宽字体，b 用加粗。"""
    for text, kind in spans:
        run = p.add_run(text)
        if kind == "code":
            _set_font(run, size=size - 0.5, color=RGBColor(0xA1, 0x3A, 0x2E),
                      bold=bold, mono=True)
        elif kind == "b":
            _set_font(run, size=size, color=color, bold=True)
        else:
            _set_font(run, size=size, color=color, bold=bold)


# ---------- 版式块 ----------
def header_block(doc, kicker: str, title: str):
    """每节顶部：一行等宽小字（文件位置）+ 大标题。"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(10)
    pf.space_after = Pt(2)
    _keep_with_next(p)
    run = p.add_run(kicker)
    _set_font(run, size=8.5, color=MUTED, mono=True)

    t = doc.add_paragraph()
    tf = t.paragraph_format
    tf.space_before = Pt(0)
    tf.space_after = Pt(8)
    tf.line_spacing = 1.25
    _keep_with_next(t)
    run = t.add_run(title)
    _set_font(run, size=15, color=INK, bold=True)


def code_block(doc, lines):
    """带语法着色的代码块：浅灰底 + 左侧细线，固定行高保证不挤。"""
    total = len(lines)
    for i, line in enumerate(lines):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(0.35)
        pf.right_indent = Cm(0.15)
        pf.space_before = Pt(4) if i == 0 else Pt(0)
        pf.space_after = Pt(4) if i == total - 1 else Pt(0)
        pf.line_spacing = Pt(13.8)          # 固定行高：代码不会因行距倍数被撑开
        if i < total - 1:                   # 代码块整体不许被分页劈开
            _keep_with_next(p)
        _shade(p._p.get_or_add_pPr(), "F4F6F9")
        _para_border(p, "left", "A9B7C9", 12, 6)
        if not line:
            _set_font(p.add_run(" "), size=9.5, color=CODE_COLORS["p"][0], mono=True)
            continue
        for text, kind in line:
            color, bold = CODE_COLORS.get(kind, CODE_COLORS["p"])
            run = p.add_run(text)
            _set_font(run, size=9.5, color=color, bold=bold, mono=True)
            if kind == "c":
                run.font.italic = True


def note_card(doc, note):
    """一条批注：竖线颜色表示语气（plain 灰 / teal 结论 / warn 知识点）。"""
    color, line_hex = TONE_COLORS.get(note.get("tone", "plain"), TONE_COLORS["plain"])

    title_p = doc.add_paragraph()
    tf = title_p.paragraph_format
    tf.left_indent = Cm(0.55)
    tf.space_before = Pt(6)
    tf.space_after = Pt(1)
    _para_border(title_p, "left", line_hex, 18, 8)
    _keep_with_next(title_p)
    add_spans(title_p, note["title"], size=10.5, color=color, bold=True)

    body_p = doc.add_paragraph()
    bf = body_p.paragraph_format
    bf.left_indent = Cm(0.55)
    bf.space_before = Pt(0)
    bf.space_after = Pt(2)
    bf.line_spacing = 1.4
    _para_border(body_p, "left", line_hex, 18, 8)
    add_spans(body_p, note["body"], size=10)


def takeaway(doc, spans):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.35)
    pf.right_indent = Cm(0.1)
    pf.space_before = Pt(8)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.4
    _shade(p._p.get_or_add_pPr(), "F2F6FD")
    _para_border(p, "left", "1A4FA8", 24, 8)
    add_spans(p, spans, size=10.5, color=RGBColor(0x14, 0x2B, 0x52))


def rule(doc, before=10, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    _para_border(p, "bottom", "E3E8EF", 6, 1)


def page_number_footer(section, text: str):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"{text}　·　")
    _set_font(run, size=8.5, color=MUTED)

    # 页码域：PAGE
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    num = OxmlElement("w:t")
    num.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for el in (fld_begin, instr, fld_sep, num, fld_end):
        run._r.append(el)
    _set_font(run, size=8.5, color=MUTED)


def build() -> Path:
    pages = json.loads(SRC.read_text(encoding="utf-8"))
    doc = Document()

    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.2)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)

    normal = doc.styles["Normal"]
    normal.font.name = CJK
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)
    page_number_footer(section, "代码解读 · 外刊精读教练")

    # ---------- 封面 ----------
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(60)
    p.paragraph_format.space_after = Pt(6)
    _set_font(p.add_run("期末项目 · 代码拆解"), size=10, color=ACCENT, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    _set_font(p.add_run("代码解读"), size=34, color=INK, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    _set_font(p.add_run("十段关键代码：配置层 / 编排层 / 工具层 / 界面层"),
              size=13, color=MUTED)

    rule(doc, before=2, after=16)

    intro = [
        "这份文档是《代码解读.pptx》（13 页）的文字版：每节保留一段真实代码、四条批注和一句结论，"
        "代码由脚本从幻灯片解析而来，与项目源码逐字一致。",
        "十段代码不是写完再挑出来讲的，而是项目跑起来后被 bug 逼出来的十次判断。"
        "它们没有一段是让模型更聪明，做的都是同一件事——把不该由模型决定的事情，从模型手里拿走。",
    ]
    for text in intro:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(8)
        pf.line_spacing = 1.55
        _set_font(p.add_run(text), size=10.5, color=RGBColor(0x33, 0x3D, 0x4B))

    # ---------- 项目信息（填充封面下半页） ----------
    info = [
        ("技术栈", "LangGraph · 通义千问 qwen-plus · Streamlit"),
        ("构成", "7 个工具 · 5 个语种 · 6161 条可核验词库 · 22 项自检"),
        ("日期", "2026 年 10 月"),
    ]
    for i, (label, value) in enumerate(info):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(26 if i == 0 else 4)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.4
        _set_font(p.add_run(f"{label}　"), size=9, color=MUTED, bold=True)
        _set_font(p.add_run(value), size=9, color=RGBColor(0x33, 0x3D, 0x4B))

    # ---------- 目录 ----------
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(12)
    _set_font(p.add_run("目录"), size=18, color=INK, bold=True)

    code_pages = [pg for pg in pages if pg.get("code")]
    for pg in code_pages:
        kicker = pg["kicker"]                      # 形如「代码 01 · config.py:40–72 · 配置层」
        head = kicker.split("·")[0].strip()        # 代码 01
        where = "·".join(kicker.split("·")[1:]).strip()
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(0.2)
        pf.space_before = Pt(4)
        pf.space_after = Pt(4)
        pf.line_spacing = 1.3
        _set_font(p.add_run(f"{head}　"), size=10.5, color=ACCENT, bold=True)
        _set_font(p.add_run(pg["h2"]), size=10.5, color=INK)
        _set_font(p.add_run(f"　　{where}"), size=8.5, color=MUTED, mono=True)

    # ---------- 使用说明（顺带把目录页下半页填满） ----------
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(34)
    p.paragraph_format.space_after = Pt(4)
    _set_font(p.add_run("怎么用这份文档"), size=11, color=INK, bold=True)

    for text in [
        "对照演示看：每节对应《代码解读.pptx》的一页，代码和批注逐字一致，讲的时候手指屏幕就行。",
        "只想快速回顾：每节最后那一条「一句话」就是结论，十条连着读完大约两分钟。",
        "被问到细节：每条批注的小标题，基本就是老师可能追问的那个问题，答案在下一行。",
    ]:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(0.75)
        pf.first_line_indent = Cm(-0.35)
        pf.space_before = Pt(2)
        pf.space_after = Pt(2)
        pf.line_spacing = 1.5
        _set_font(p.add_run("· "), size=10, color=ACCENT, bold=True)
        _set_font(p.add_run(text), size=10, color=RGBColor(0x33, 0x3D, 0x4B))

    # ---------- 正文（每段独立成页，代码块不会被劈开） ----------
    for i, pg in enumerate(code_pages):
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        header_block(doc, pg["kicker"], pg["h2"])
        code_block(doc, pg["code"])
        for note in pg.get("notes", []):
            note_card(doc, note)
        if pg.get("takeaway"):
            takeaway(doc, pg["takeaway"])

    # ---------- 收束 ----------
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    _set_font(p.add_run("十段代码，共用一条判断"), size=15, color=INK, bold=True)
    for pg in pages:
        if pg.get("takeaway") and not pg.get("code"):
            takeaway(doc, pg["takeaway"])

    # 复习用：十段结论一览
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    _set_font(p.add_run("十段结论一览"), size=11, color=INK, bold=True)

    for pg in code_pages:
        num = pg["kicker"].split("·")[0].replace("代码", "").strip()
        text = spans_text(pg["takeaway"]).replace("一句话：", "", 1).strip()
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(1.1)
        pf.first_line_indent = Cm(-1.1)
        pf.space_before = Pt(5)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.45
        _set_font(p.add_run(f"{num}　"), size=9, color=ACCENT, bold=True, mono=True)
        _set_font(p.add_run(text), size=9.5, color=RGBColor(0x33, 0x3D, 0x4B))

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    out = build()
    print(f"已生成 {out}  ({out.stat().st_size / 1024:.0f} KB)")
